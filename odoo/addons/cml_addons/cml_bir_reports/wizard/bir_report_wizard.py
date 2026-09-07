from odoo import models, fields, api
import os
import openpyxl
from io import BytesIO
import base64
from datetime import datetime, date
from dateutil.relativedelta import relativedelta
from openpyxl.utils import get_column_letter
from calendar import monthrange


class BirReportWizard(models.TransientModel):
    _name = 'bir.report.wizard'
    _description = 'BIR Report Wizard'

    template_id = fields.Many2one('bir.report.template', string="Report Template", required=True)

    def get_quarter(self, date_value):
        month = date_value.month
        return (month - 1) // 3 + 1

    def get_first_day_of_the_quarter(self, idate, quarter):
        return date(idate.year, 3 * quarter - 2, 1)

    def write_currency(self, ws, cell, value):
        """Write number as currency (comma separated with 2 decimals)."""
        ws[cell] = float(value or 0.0)
        ws[cell].number_format = '#,##0.00'

    def action_generate_excel_report(self):
        active_id = self.env.context.get('active_id')
        if not active_id:
            raise ValueError("No active invoice found!")

        wb = openpyxl.load_workbook(filename=BytesIO(base64.b64decode(self.template_id.file)))
        ws = wb.active
        invoice = self.env['account.invoice'].browse(active_id)

        # Payee (Customer) Info
        ws['B20'] = invoice.partner_id.name or 'N/A'
        ws['B23'] = invoice.partner_id.street + " " + invoice.partner_id.street2 + ", " + invoice.partner_id.city + ", " + invoice.partner_id.state_id.name + ", " + invoice.partner_id.country_id.name

        zip_code = invoice.partner_id.zip
        zip_cells = ['AK24', 'AL24', 'AM24', 'AN24']
        if zip_code:
            for i in range(min(len(zip_code), len(zip_cells))):
                ws[zip_cells[i]] = zip_code[i]

        vat_number = invoice.partner_id.vat
        vat_cells = ['N17', 'O17', 'P17', 'R17', 'S17', 'T17', 'V17', 'W17', 'X17', 'Z17', 'AA17', 'AB17', 'AC17', 'AD17']
        if vat_number:
            for i in range(min(len(vat_number), len(vat_cells))):
                ws[vat_cells[i]] = vat_number[i]

        # Payor (Company) Info
        ws['B35'] = invoice.company_id.name or 'N/A'
        ws['B38'] = invoice.company_id.street + ", " + invoice.company_id.city + ", " + invoice.company_id.state_id.name + ", " + invoice.company_id.country_id.name

        comp_zip = invoice.company_id.zip
        comp_zip_cells = ['AK39', 'AL39', 'AM39', 'AN39']
        if comp_zip:
            for i in range(min(len(comp_zip), len(comp_zip_cells))):
                ws[comp_zip_cells[i]] = comp_zip[i]

        comp_vat = invoice.company_id.vat
        comp_vat_cells = ['N32', 'O32', 'P32', 'R32', 'S32', 'T32', 'V32', 'W32', 'X32', 'Z32', 'AA32', 'AB32', 'AC32', 'AD32', 'AE32']
        if comp_vat:
            for i in range(min(len(comp_vat), len(comp_vat_cells))):
                ws[comp_vat_cells[i]] = comp_vat[i]

        # Dates
        current_date = invoice.date_invoice
        quarter = self.get_quarter(current_date)
        sdate = self.get_first_day_of_the_quarter(current_date, quarter)
        edate = (sdate + relativedelta(months=3, seconds=-1)).date()

        s_date = sdate.strftime('%m%d%Y')
        e_date = edate.strftime('%m%d%Y')

        if s_date:
            for i in range(len(s_date)):
                col_letter = get_column_letter(10 + i)
                ws[f'{col_letter}12'] = s_date[i]
        if e_date:
            for i in range(len(e_date)):
                col_letter = get_column_letter(27 + i)
                ws[f'{col_letter}12'] = e_date[i]

        # Footer
        ws['A70'] = self.template_id.company_profile_name + '\n' + self.template_id.company_position_name

        # Taxes
        total_amount = 0
        total_tax_base = 0
        row = 45

        for tax in invoice.tax_line_ids:
            if not tax.tax_id.show_in_tax_report:
                continue

            ws[f'A{row}'] = tax.description
            ws[f'L{row}'] = tax.name
            self.write_currency(ws, f'AI{row}', abs(tax.amount_total))
            total_amount += abs(tax.amount_total)

            tax_base = tax.base
            if current_date.month in [1, 4, 7, 10]:
                self.write_currency(ws, f'O{row}', tax_base)
                self.write_currency(ws, f'AD{row}', tax_base)
            elif current_date.month in [2, 5, 8, 11]:
                self.write_currency(ws, f'T{row}', tax_base)
                self.write_currency(ws, f'AD{row}', tax_base)
            elif current_date.month in [3, 6, 9, 12]:
                self.write_currency(ws, f'Y{row}', tax_base)
                self.write_currency(ws, f'AD{row}', tax_base)

            total_tax_base += tax_base
            row += 1

        # Totals
        if current_date.month in [1, 4, 7, 10]:
            self.write_currency(ws, 'O55', total_tax_base)
            self.write_currency(ws, 'AD55', total_tax_base)
        elif current_date.month in [2, 5, 8, 11]:
            self.write_currency(ws, 'T55', total_tax_base)
            self.write_currency(ws, 'AD55', total_tax_base)
        elif current_date.month in [3, 6, 9, 12]:
            self.write_currency(ws, 'Y55', total_tax_base)
            self.write_currency(ws, 'AD55', total_tax_base)

        self.write_currency(ws, 'AI55', total_amount)

        # Save output
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        file_data = base64.b64encode(output.getvalue())

        attachment = self.env['ir.attachment'].create({
            'name': f'{invoice.number}.xlsx',
            'type': 'binary',
            'datas': file_data,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',

            'datas_fname': f'{invoice.number}.xlsx',
            'res_model': 'account.invoice',
            'res_id': invoice.id,
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }

